"""Word version of the pilot sheet, and the way back to labels CSV.
  python scripts/llm/pilot_doc.py make                      -> data/llm/pilot/pilot_labelling.docx (150 items, blank Label column)
  python scripts/llm/pilot_doc.py read --rater kareem FILE  -> data/llm/pilot/labels_kareem.csv (a/n/d or full words accepted)
Only parent and reply text are shown (no usernames, week, subreddit, stratum or model output)."""
import argparse
import csv
import sys
from pathlib import Path

import docx
from docx.shared import Cm, Pt

P = Path("data/llm/pilot")
FULL = {"a": "agree", "agree": "agree", "n": "neutral", "neutral": "neutral", "d": "disagree", "disagree": "disagree"}

ap = argparse.ArgumentParser()
ap.add_argument("mode", choices=["make", "read", "make-md", "read-md"])
ap.add_argument("file", nargs="?")
ap.add_argument("--rater")
ap.add_argument("--force", action="store_true")
a = ap.parse_args()

import re


def quote(t):
    return "\n".join("> " + l for l in t.splitlines()) or ">"


if a.mode == "make-md":
    rows = list(csv.DictReader(open(P / "pilot_sheet_blank.csv", newline="")))
    out = ["# Pilot labelling sheet\n",
           "Type agree, neutral or disagree (or a / n / d) after **Label:** for each item. Question: what is the REPLY's stance "
           "toward the position taken in the PARENT comment? Tone is not stance. Judge only the text shown. Label independently. "
           "See docs/pilot_rubric.md.\n"]
    for r in rows:
        out.append(f"---\n\n### {r['pilot_id']}\n\n**Parent:**\n{quote(r['parent_comment'])}\n\n**Reply:**\n{quote(r['reply'])}\n\n**Label:** \n")
    (P / "pilot_labelling.md").write_text("\n".join(out))
    print(f"wrote {P / 'pilot_labelling.md'} ({len(rows)} items)")
elif a.mode == "read-md":
    if not (a.file and a.rater):
        sys.exit("usage: read-md --rater NAME FILE.md")
    out = P / f"labels_{a.rater}.csv"
    if out.exists() and not a.force:
        sys.exit(f"{out} exists (labels from the terminal tool?). Use --force to overwrite.")
    labels, bad = [], []
    for pid, body in re.findall(r"### (P\d+)\n(.*?)(?=\n---|\Z)", Path(a.file).read_text(), re.S):
        raw = re.search(r"\*\*Label:\*\*[ \t]*([^\n]*)", body).group(1).strip().lower().strip(".")
        if raw in FULL:
            labels.append((pid, FULL[raw]))
        elif raw:
            bad.append((pid, raw))
    with open(out, "w", newline="") as f:
        w = csv.writer(f)
        w.writerow(["pilot_id", "label"])
        w.writerows(labels)
    print(f"wrote {len(labels)} labels to {out}; unreadable: {bad}")
elif a.mode == "make":
    rows = list(csv.DictReader(open(P / "pilot_sheet_blank.csv", newline="")))
    d = docx.Document()
    d.styles["Normal"].font.size = Pt(10)
    for s in d.sections:
        s.left_margin = s.right_margin = Cm(1.8)
    d.add_heading("Pilot labelling sheet", 1)
    d.add_paragraph("For each item, type one word in the Label column: agree, neutral or disagree (a / n / d also works). "
                    "Question: what is the REPLY's stance toward the position taken in the PARENT comment? Judge only the text shown; "
                    "tone is not stance. Label independently; do not compare with the other rater until both are done. "
                    "See docs/pilot_rubric.md for examples and the hard cases.")
    t = d.add_table(rows=1, cols=3)
    t.style = "Table Grid"
    for c, h in zip(t.rows[0].cells, ["Item", "Parent comment, then reply", "Label"]):
        c.text = h
        c.paragraphs[0].runs[0].bold = True
    for r in rows:
        c = t.add_row().cells
        c[0].text = r["pilot_id"]
        p = c[1].paragraphs[0]
        p.add_run("PARENT: ").bold = True
        p.add_run(r["parent_comment"])
        q = c[1].add_paragraph()
        q.add_run("REPLY: ").bold = True
        q.add_run(r["reply"])
    for row in t.rows:
        for c, w in zip(row.cells, [Cm(1.5), Cm(13.5), Cm(2.5)]):
            c.width = w
    d.save(P / "pilot_labelling.docx")
    print(f"wrote {P / 'pilot_labelling.docx'} ({len(rows)} items)")
else:
    if not (a.file and a.rater):
        sys.exit("usage: read --rater NAME FILE.docx")
    out = P / f"labels_{a.rater}.csv"
    if out.exists() and not a.force:
        sys.exit(f"{out} exists (labels from the terminal tool?). Use --force to overwrite.")
    t = docx.Document(a.file).tables[0]
    ids = {r["pilot_id"] for r in csv.DictReader(open(P / "pilot_sheet_blank.csv", newline=""))}
    labels, bad = [], []
    for row in t.rows[1:]:
        pid, raw = row.cells[0].text.strip(), row.cells[2].text.strip().lower().strip(".")
        if pid not in ids:
            bad.append((pid, "unknown id"))
        elif raw == "":
            continue
        elif raw in FULL:
            labels.append((pid, FULL[raw]))
        else:
            bad.append((pid, raw))
    with open(out, "w", newline="") as f:
        w = csv.writer(f)
        w.writerow(["pilot_id", "label"])
        w.writerows(labels)
    print(f"wrote {len(labels)} labels to {out}; blank: {len(ids) - len(labels) - len(bad)}; unreadable: {bad}")
